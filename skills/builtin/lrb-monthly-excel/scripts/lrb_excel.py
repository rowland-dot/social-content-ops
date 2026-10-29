#!/usr/bin/env python3
"""LRB 月度计划 / 月报 Excel 的唯一做法（issue 57）。

    build  团队底稿进，客户文件出：整理形状 → 套用负责人批准的内容改动 → 调用负责人的 LRB 文件夹里的
           planstyle.py / reportstyle.py 排版 → 跑检查。
    check  对任何一份已排好版的文件跑同一组检查，不改文件（用来核对已有的成品）。

不改、不复制、不分叉 LRB 文件夹里的脚本和 verify()；只在原地导入它们。导入时不写字节码，
不在负责人的文件夹里留任何文件。排版步骤只改格式，不改任何标题、数字、日期、链接和评论文字；
内容改动只来自 --changes 里负责人批准的条目，并单独列出。

退出码：0 全部通过；1 排版或检查没过（文件不能交）；2 跑不起来（参数、路径、底稿形状、LRB 文件夹）。
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True  # 导入负责人的脚本时不能往他的文件夹写 __pycache__

SHARED_SCRIPTS = Path(__file__).resolve().parents[3] / "shared" / "scripts"
def _default_lrb() -> tuple[str, ...]:
    """The LRB folder sits in the Windows user's folder; find it there without naming the user."""
    tail = "Claude Cowork/LRB Plan & Report"
    users = Path("/mnt/c/Users")
    found = sorted(str(p) for p in users.glob(f"*/{tail}")) if users.is_dir() else []
    return (*found, str(Path.home() / tail))


DEFAULT_LRB = _default_lrb()

IDEO = re.compile(r"[\u3400-\u9fff]")  # 汉字（和 lrbstyle 的 cjk 里的汉字区一致）
LATIN = re.compile(r"[A-Za-z]")
WORD = re.compile(r"[A-Za-z][A-Za-z'’-]*")
BAD_CHARS = re.compile("[\u2010\u2011\u2012\u200d\u00a0]")  # 剧本第 4 节：字符扫描必须为 0
LOGO_MAX_H, LOGO_MAX_W, PROPORTION_TOLERANCE = 84, 280, 0.03

# 排版脚本把全角字符算成中文，所以英文行里的全角字符（｜ ！ ） 等）会让两行并成一行。
# 排版前把它们换成私有区字符，排版后原样换回：文字一个字符都不变。
FULLWIDTH_RANGES = ((0xFF01, 0xFF5E, 0xE000), (0x3001, 0x303F, 0xE100))


class Cannot(Exception):
    """跑不起来（退出码 2）。"""


# ---- 全角保护 ----------------------------------------------------------------------------------


def to_private(ch: str) -> str:
    for lo, hi, base in FULLWIDTH_RANGES:
        if lo <= ord(ch) <= hi:
            return chr(base + ord(ch) - lo)
    return ch


def is_fullwidth(ch: str) -> bool:
    return to_private(ch) != ch


def private_to_original() -> dict[str, str]:
    return {to_private(chr(c)): chr(c) for lo, hi, _ in FULLWIDTH_RANGES for c in range(lo, hi + 1)}


def is_english_line(line: str) -> bool:
    return not IDEO.search(line) and bool(LATIN.search(line))


def ascii_offer(line: str) -> str:
    """建议的半角写法（内容改动，只列出来，不执行）。"""
    line = re.sub(r"\s*｜\s*", ": ", line)
    return "".join(unicodedata.normalize("NFKC", ch) if 0xFF01 <= ord(ch) <= 0xFF5E else ch for ch in line)


def protect(ws) -> dict[str, list[str]]:
    """把英文行里的全角字符换成私有区字符；返回 {单元格: [被保护的英文行]}。"""
    kept: dict[str, list[str]] = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith("="):
                continue
            if any(0xE000 <= ord(ch) <= 0xE1FF for ch in v):
                raise Cannot(f"{c.coordinate} already holds private-use characters; the draft cannot be protected")
            lines, hit = v.split("\n"), []
            for i, line in enumerate(lines):
                if is_english_line(line) and any(is_fullwidth(ch) for ch in line):
                    hit.append(line.strip())
                    lines[i] = "".join(to_private(ch) for ch in line)
            if hit:
                c.value = "\n".join(lines)
                kept[c.coordinate] = hit
    return kept


def restore(path: Path) -> None:
    """成品里的私有区字符原样换回全角；只动字符串部分，别的字节不碰。"""
    back = private_to_original()

    def put_back(m):  # openpyxl 写成 &#57435; 这样的数字引用，Excel 写成原字符：两种都认
        code = int(m.group(1)[1:], 16) if m.group(1)[0] in "xX" else int(m.group(1))
        return back.get(chr(code), m.group(0))

    tmp = str(path) + ".restore"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/sharedStrings.xml" or re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                text = re.sub(r"&#(x[0-9A-Fa-f]+|\d+);", put_back, data.decode("utf-8"))
                data = "".join(back.get(ch, ch) for ch in text).encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


# ---- 整理形状（只动格式，不动文字） -----------------------------------------------------------------


def _merges(ws):
    from openpyxl.worksheet.cell_range import CellRange

    ranges = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in ranges:
        ws.unmerge_cells(str(m))
    return ranges


def delete_col(ws, idx: int) -> None:
    merges = _merges(ws)
    ws.delete_cols(idx)
    for m in merges:
        if m.min_col == m.max_col == idx:
            continue
        lo = m.min_col - (1 if m.min_col > idx else 0)
        hi = m.max_col - (1 if m.max_col >= idx else 0)
        ws.merge_cells(start_row=m.min_row, start_column=lo, end_row=m.max_row, end_column=hi)


def delete_row(ws, idx: int) -> None:
    merges = _merges(ws)
    old_max = ws.max_row
    heights = {r: ws.row_dimensions[r].height for r in range(1, old_max + 2)}
    ws.delete_rows(idx)
    for r in range(idx, old_max + 1):  # 行高跟着上移，最后一行不留旧行高
        ws.row_dimensions[r].height = heights.get(r + 1)
    for m in merges:
        if m.min_row == m.max_row == idx:
            continue
        lo = m.min_row - (1 if m.min_row > idx else 0)
        hi = m.max_row - (1 if m.max_row >= idx else 0)
        ws.merge_cells(start_row=lo, start_column=m.min_col, end_row=hi, end_column=m.max_col)


def work_rows(ws, col: int) -> list[int]:
    return [r for r in range(3, ws.max_row + 1) if "Work" in str(ws.cell(r, col).value or "")]


def row_is_empty(ws, r: int) -> bool:
    return all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1))


def normalise_plan(ws) -> list[str]:
    """排版脚本要求：表从 A 列开始，只有一个 Work 行，日程和 Work 之间正好一行空行。"""
    from openpyxl.utils import get_column_letter

    notes = []
    first = min((c.column for row in ws.iter_rows() for c in row if c.value is not None), default=1)
    for _ in range(first - 1):
        delete_col(ws, 1)
    if first > 1:
        notes.append(f"moved the table from column {get_column_letter(first)} to column A")
    rows = work_rows(ws, 1)
    if not rows:
        raise Cannot("the draft has no row with 'Work' in its first column, so the work plan block cannot be found")
    labelled = [r for r in rows if ws.cell(r, 2).value not in (None, "")]
    if len(rows) > 1 and labelled:
        for r in sorted((x for x in rows if x < max(labelled) and x not in labelled), reverse=True):
            text = str(ws.cell(r, 1).value)
            delete_row(ws, r)
            notes.append(f"removed duplicate work-plan heading row {r}: '{text}'")
    wp = work_rows(ws, 1)[0]
    last = max((r for r in range(3, wp) if ws.cell(r, 1).value is not None), default=2)
    removed = 0
    for r in range(wp - 1, last + 1, -1):
        if row_is_empty(ws, r):
            delete_row(ws, r)
            removed += 1
    if removed:
        notes.append(f"removed {removed} extra blank row{'s' if removed > 1 else ''} between the schedule and the work plan")
    return notes


# ---- 负责人批准的内容改动 --------------------------------------------------------------------------


QUOTE_PAIRS = (('"', '"'), ("“", "”"), ("「", "」"))  # " "  “ ”  「 」
CELL_ADDRESS = re.compile(r"[A-Za-z]{1,3}[1-9][0-9]*")
DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d")


def has_quoted_words(ruling: str) -> bool:
    """A ruling is the owner's own words in quotation marks, not a summary of them (issue 57, round 3)."""
    return any(m.group(1).strip() for a, b in QUOTE_PAIRS
               for m in re.finditer(re.escape(a) + r"(.*?)" + re.escape(b), ruling, re.S))


def load_changes(path: str) -> list[dict]:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise Cannot(f"cannot read the content changes file {path}: {exc}")
    if not isinstance(data, list):
        raise Cannot("the content changes file must hold a list of {old, new, ruling}")
    for i, ch in enumerate(data, 1):
        for key in ("old", "new", "ruling"):
            if not isinstance(ch, dict) or not isinstance(ch.get(key), str) or not ch[key].strip():
                raise Cannot(f"content change {i} needs a non-empty '{key}' (the ruling is the owner's own words)")
        if not has_quoted_words(ch["ruling"]):
            raise Cannot(f"content change {i}: the ruling must quote the owner's own words in quotation marks, not summarise "
                         f"them (got: {ch['ruling']!r})")
        if "cell" in ch and ("date" in ch or "column" in ch):
            raise Cannot(f"content change {i}: name the cell with either \"cell\" or \"date\" and \"column\", not both")
        if "date" in ch and not isinstance(ch.get("column"), str):
            raise Cannot(f"content change {i}: a \"date\" needs a \"column\" (a letter in the finished file)")
        if "column" in ch and "date" not in ch:
            raise Cannot(f"content change {i}: a \"column\" only goes with a \"date\"")
    return data


def parse_date(value):
    import datetime

    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def draft_column(letters: str, kind: str, i: int) -> int:
    """A column letter of the finished file -> the draft's column number (a plan's finished file adds a gutter)."""
    from openpyxl.utils import column_index_from_string

    if not re.fullmatch(r"[A-Za-z]{1,3}", letters):
        raise Cannot(f"content change {i}: {letters!r} is not a column letter")
    col = column_index_from_string(letters.upper()) - (1 if kind == "plan" else 0)
    if col < 1:
        raise Cannot(f"content change {i}: column {letters.upper()} is the gutter column of the finished file; it holds no text")
    return col


def locate(ws, i: int, ch: dict, kind: str):
    """The one cell a change names, or None when it names only text. Addresses are the finished file's own."""
    from openpyxl.utils import get_column_letter

    if "cell" in ch:
        if not isinstance(ch["cell"], str) or not CELL_ADDRESS.fullmatch(ch["cell"].strip()):
            raise Cannot(f"content change {i}: {ch['cell']!r} is not a cell address (use one like C10)")
        letters = re.match(r"[A-Za-z]+", ch["cell"].strip()).group().upper()
        row = int(ch["cell"].strip()[len(letters):])
        if kind == "plan" and letters == "A":
            raise Cannot(f"content change {i}: cell {ch['cell'].strip()} is in the gutter column of the finished file; it holds no text")
        return ws.cell(row, draft_column(letters, kind, i))
    if "date" in ch:
        want = parse_date(ch["date"])
        if want is None:
            raise Cannot(f"content change {i}: {ch['date']!r} is not a date (use 2026-11-27 or 27/11/2026)")
        col = draft_column(ch["column"], kind, i)
        rows = [r for r in range(3, ws.max_row + 1) if parse_date(ws.cell(r, 1).value) == want]
        if len(rows) != 1:
            raise Cannot(f"content change {i}: {len(rows)} schedule rows have the date {ch['date']} "
                         f"(it must be exactly one)")
        return ws.cell(rows[0], col)
    return None


def apply_changes(ws, changes: list[dict], kind: str = "plan") -> list[tuple[str, dict]]:
    applied = []
    for i, ch in enumerate(changes, 1):
        old = ch["old"].replace("\r", "")
        cell = locate(ws, i, ch, kind)
        if cell is not None:
            held = cell.value if isinstance(cell.value, str) else ""
            if held.replace("\r", "") != old:
                raise Cannot(f"content change {i}: cell {out_coord(cell.coordinate, kind)} holds {held!r}, "
                             f"not the old text {ch['old']!r}")
            hits = [cell]
        else:
            hits = [c for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.replace("\r", "") == old]
        if len(hits) != 1:
            names = ", ".join(out_coord(c.coordinate, kind) for c in hits)
            raise Cannot(f"content change {i}: the old text matches {len(hits)} cells in the draft"
                         + (f" ({names}); give \"cell\" with the one to change, or \"date\" and \"column\"" if hits else "")
                         + "; the old text must be the whole text of the cell, copied exactly")
        hits[0].value = ch["new"]
        applied.append((hits[0].coordinate, ch))
    return applied


def out_coord(coord: str, kind: str) -> str:
    """底稿坐标 → 成品坐标：计划排版时左边多出一列空白边。"""
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.utils.cell import coordinate_from_string

    letters, row = coordinate_from_string(coord)
    shift = 1 if kind == "plan" else 0
    return f"{get_column_letter(column_index_from_string(letters) + shift)}{row}"


def cell_text(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>")


def write_changes_file(path: Path, applied, offered) -> None:
    lines = ["# 内容改动清单（单独列出，不在排版里）", "", "## 已执行（负责人已批准）", ""]
    if applied:
        lines += ["| 单元格 | 原文 | 新文 | 负责人的裁定 |", "|---|---|---|---|"]
        lines += [f"| {c} | {cell_text(ch['old'])} | {cell_text(ch['new'])} | {cell_text(ch['ruling'])} |" for c, ch in applied]
    else:
        lines.append("无。")
    lines += ["", "## 建议，等负责人裁定（没有执行）", ""]
    if offered:
        lines += ["| 单元格 | 原文 | 建议 |", "|---|---|---|"]
        lines += [f"| {c} | {cell_text(old)} | {cell_text(new)} |" for c, old, new in offered]
    else:
        lines.append("无。")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---- 检查 ---------------------------------------------------------------------------------------


def check_verify(path: Path, lrbstyle) -> str | None:
    try:
        lrbstyle.verify(str(path))
    except SystemExit as exc:
        return str(exc.code)
    except Exception as exc:  # noqa: BLE001 — verify 本身出错也是一次没过的检查
        return f"{type(exc).__name__}: {exc}"
    return None


def check_logo(ws) -> str | None:
    logos = [im for im in ws._images if getattr(getattr(im.anchor, "_from", None), "row", None) == 0]
    if not logos:
        return "no logo in the title row (give --logo when the draft has none)"
    bad = []
    for im in logos:
        ext = getattr(im.anchor, "ext", None)
        if ext is None:
            bad.append("logo was not placed by the formatting step (it has no fixed size)")
            continue
        w, h = ext.width / 9525, ext.height / 9525
        if round(h) > LOGO_MAX_H:
            bad.append(f"logo is {round(h)} px tall, the playbook caps it at {LOGO_MAX_H}")
        if round(w) > LOGO_MAX_W:
            bad.append(f"logo is {round(w)} px wide, the playbook caps it at {LOGO_MAX_W}")
        if abs((w / h) / (im.width / im.height) - 1) > PROPORTION_TOLERANCE:
            bad.append(f"logo is {round(w)}x{round(h)} but its picture is {im.width}x{im.height}: proportions not kept")
    return "; ".join(bad) or None


def check_spacer(ws) -> str | None:
    rows = work_rows(ws, 2)
    if not rows:
        return "no row with 'Work' in column B"
    wp = rows[0]
    last = max((r for r in range(3, wp) if ws.cell(r, 2).value is not None), default=2)
    n = wp - last - 1
    return None if n == 1 else f"{n} spacer rows between the schedule and the work plan, the playbook has exactly 1"


def check_bilingual(ws, kind: str) -> str | None:
    wp = (work_rows(ws, 2) or [None])[0] if kind == "plan" else None
    bad = []
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or (wp and c.row >= wp and c.column >= 4):
                continue
            lines = [ln for ln in c.value.replace("\r", "").split("\n") if ln.strip()]
            if any(IDEO.search(ln) and len(WORD.findall(ln)) >= 3 for ln in lines):
                bad.append(f"{c.coordinate} (Chinese and English share a line)")
            elif len(lines) >= 2 and not IDEO.search(lines[0]) and any(IDEO.search(ln) for ln in lines[1:]):
                bad.append(f"{c.coordinate} (English above Chinese)")
    return "; ".join(bad) or None


def check_chars(ws) -> str | None:
    bad = [f"{c.coordinate} U+{ord(m.group()):04X}" for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) for m in [BAD_CHARS.search(c.value)] if m]
    return ("stray characters: " + ", ".join(bad[:8])) if bad else None


def check_font(ws) -> str | None:
    bad = [f"{c.coordinate} {c.font.name}" for row in ws.iter_rows() for c in row
           if c.value is not None and c.font.name != "Arial"]
    return ("not Arial: " + ", ".join(bad[:8]) + (f" (+{len(bad) - 8} more)" if len(bad) > 8 else "")) if bad else None


def fill_groups(ws) -> dict[str, list[str]]:
    """Every cell that carries any fill (white, theme, empty-but-styled ones too), grouped by colour."""
    groups: dict[str, list[str]] = {}
    for row in ws.iter_rows():
        for c in row:
            f = c.fill
            if f is None or getattr(f, "fill_type", "gradient") is None:  # a gradient has no fill_type
                continue
            col = getattr(f, "fgColor", None)
            if col is None:
                key = "gradient"
            elif col.type == "rgb" and isinstance(col.rgb, str):
                key = col.rgb[-6:]
            else:
                key = f"theme{col.theme}" if col.type == "theme" else f"indexed{col.indexed}"
            groups.setdefault(key, []).append(c.coordinate)
    return groups


def describe_groups(groups: dict[str, list[str]], coord=lambda c: c) -> list[str]:
    def where(cells: list[str]) -> str:
        from openpyxl.utils import column_index_from_string as idx, get_column_letter as letter
        from openpyxl.utils.cell import coordinate_from_string

        shown = [coord(c) for c in cells]
        if len(shown) <= 6:
            return ",".join(shown)
        parts = [coordinate_from_string(c) for c in shown]  # many cells: the box that holds them all
        cols, rows = [idx(p[0]) for p in parts], [p[1] for p in parts]
        return f"in the box {letter(min(cols))}{min(rows)}:{letter(max(cols))}{max(rows)}"

    return [f"{k} {len(v)} cell{'s' if len(v) > 1 else ''} {where(v)}" for k, v in groups.items()]


def clear_fills(ws) -> dict[str, list[str]]:
    """The owner, 2026-10-17: "No colours, per template". The playbook's look has no fills at all, so any colour a
    draft carries is removed before the formatting step sees it (formatting only: no value is touched)."""
    from openpyxl.styles import PatternFill

    groups = fill_groups(ws)
    for cells in groups.values():
        for coordinate in cells:
            ws[coordinate].fill = PatternFill(fill_type=None)
    return groups


def check_fills(ws) -> str | None:
    found = describe_groups(fill_groups(ws))
    return ("colours on the cells (the playbook has no fills): " + "; ".join(found)) if found else None


def run_checks(path: Path, kind: str, lrbstyle) -> list[tuple[str, str | None]]:
    import openpyxl

    ws = openpyxl.load_workbook(path).worksheets[0]
    results = [("verify", check_verify(path, lrbstyle)), ("logo", check_logo(ws))]
    if kind == "plan":
        results.append(("spacer", check_spacer(ws)))
    results += [("bilingual", check_bilingual(ws, kind)), ("chars", check_chars(ws)), ("font", check_font(ws)),
                ("fills", check_fills(ws))]
    return results


def print_results(results) -> int:
    failed = 0
    for name, reason in results:
        if reason is None:
            print(f"PASS {name}")
        else:
            failed += 1
            print(f"FAIL {name}: {' '.join(reason.split())}")
    return failed


# ---- LRB 文件夹 / 输出路径 ------------------------------------------------------------------------


def find_lrb(arg: str | None) -> Path:
    tried = [arg] if arg else [os.environ.get("LRB_DIR"), *DEFAULT_LRB]
    for c in tried:
        if c and (Path(c) / "playbooks" / "scripts").is_dir():
            return Path(c)
    raise Cannot(f"LRB folder not found (tried: {', '.join(t for t in tried if t)}); pass --lrb-dir or set LRB_DIR")


def load_lrb(lrb: Path, names: tuple[str, ...]):
    import importlib

    scripts = str(lrb / "playbooks" / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    mods = []
    for name in names:
        try:
            mods.append(importlib.import_module(name))
        except Exception as exc:  # noqa: BLE001
            raise Cannot(f"LRB script {name}.py cannot be imported from {scripts}: {type(exc).__name__}: {exc}")
    return mods


def validate_out(value: str) -> Path:
    sys.path.insert(0, str(SHARED_SCRIPTS))
    try:
        from output_paths import validate_output_path  # 项目统一的输出路径规约

        return validate_output_path(value, create_parent=True)
    except Exception as exc:  # noqa: BLE001 — 路径不合规矩、找不到项目根，都是"跑不起来"
        raise Cannot(str(exc))


# ---- 命令 ---------------------------------------------------------------------------------------


def cmd_check(args) -> int:
    path = Path(args.file)
    if not path.is_file():
        raise Cannot(f"{path} is not a file")
    (lrbstyle,) = load_lrb(find_lrb(args.lrb_dir), ("lrbstyle",))
    if not hasattr(lrbstyle, "verify"):
        raise Cannot("the LRB folder's lrbstyle.py has no verify(), so the file cannot be verified")
    failed = print_results(run_checks(path, args.kind, lrbstyle))
    print("RESULT OK" if not failed else f"RESULT FAILED {failed}")
    return 0 if not failed else 1


def cmd_build(args) -> int:
    import openpyxl

    lrb = find_lrb(args.lrb_dir)
    source = Path(args.source)
    if not source.is_file():
        raise Cannot(f"the draft {source} is not a file")
    out = validate_out(args.out)
    changes_out = validate_out(args.changes_out) if args.changes_out else None
    changes = load_changes(args.changes) if args.changes else []
    style_name = "planstyle" if args.kind == "plan" else "reportstyle"
    lrbstyle, style = load_lrb(lrb, ("lrbstyle", style_name))
    style_fn = getattr(style, "style_plan" if args.kind == "plan" else "style_report", None)
    if style_fn is None or not hasattr(lrbstyle, "verify"):
        raise Cannot(f"the LRB folder's {style_name}.py or lrbstyle.py is missing its entry function or verify()")

    wb = openpyxl.load_workbook(source)
    ws = wb.worksheets[0]
    shape = normalise_plan(ws) if args.kind == "plan" else []
    applied = apply_changes(ws, changes, args.kind)
    cleared = describe_groups(clear_fills(ws), lambda c: out_coord(c, args.kind))
    kept = protect(ws)
    offered = [(out_coord(c, args.kind), line, ascii_offer(line)) for c, lines in kept.items() for line in lines
               if ascii_offer(line) != line]

    for note in shape:
        print(f"SHAPE {note}")
    for line in cleared:
        print(f"CLEARED {line}")
    for c, ch in applied:
        print(f"APPLIED {out_coord(c, args.kind)}: {json.dumps(ch['old'], ensure_ascii=False)} => "
              f"{json.dumps(ch['new'], ensure_ascii=False)} | {ch['ruling']}")
    for c, lines in kept.items():
        print(f"KEPT-LINE {out_coord(c, args.kind)}: {lines[0]} (English line with full-width characters stays on line 2, text unchanged)")
    for c, old, new in offered:
        print(f"OFFERED {c}: {old} => {new}")

    with tempfile.TemporaryDirectory() as tmp:
        staged = Path(tmp) / "draft.xlsx"
        wb.save(staged)
        try:
            style_fn(str(staged), str(out), args.sheet, args.logo)
        except BaseException as exc:  # noqa: BLE001 — 排版脚本用 SystemExit 报 verify 失败
            if out.exists():
                out.unlink()
            reason = exc.code if isinstance(exc, SystemExit) else f"{type(exc).__name__}: {exc}"
            print(f"FAIL format: {' '.join(str(reason).split())}")
            print("RESULT FAILED 1")
            return 1
    if kept:
        restore(out)
    if changes_out:
        write_changes_file(changes_out, [(out_coord(c, args.kind), ch) for c, ch in applied], offered)
    failed = print_results(run_checks(out, args.kind, lrbstyle))
    print("RESULT OK" if not failed else f"RESULT FAILED {failed}")
    return 0 if not failed else 1


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="LRB 月度计划 / 月报 Excel：build（底稿→成品）和 check（核对成品）")
    sub = parser.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="团队底稿 → 客户文件")
    b.add_argument("--kind", choices=("plan", "report"), required=True)
    b.add_argument("--source", required=True, help="团队的底稿（只读，不改）")
    b.add_argument("--out", required=True, help="成品路径，必须在 outputs/<主题>/ 下")
    b.add_argument("--sheet", required=True, help='工作表名，例如 "Nov 2026"')
    b.add_argument("--logo", help="品牌 logo（.png，或带 logo 的月报）；底稿自带 logo 时可省")
    b.add_argument("--changes", help="负责人批准的内容改动（JSON 列表：old、new、ruling）")
    b.add_argument("--changes-out", help="把内容改动清单另存成 Markdown")
    b.add_argument("--lrb-dir", help="负责人的 LRB 文件夹（默认 LRB_DIR 环境变量，再试负责人机器上的位置）")
    c = sub.add_parser("check", help="核对一份已排好版的文件，不改它")
    c.add_argument("--kind", choices=("plan", "report"), required=True)
    c.add_argument("file")
    c.add_argument("--lrb-dir")
    args = parser.parse_args(argv)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print(f"cannot {args.cmd}: openpyxl is not installed for {sys.executable}; run this with Easel's own Python "
              f"(.venv/bin/python in the project folder)", file=sys.stderr)
        return 2
    try:
        return cmd_build(args) if args.cmd == "build" else cmd_check(args)
    except Cannot as exc:
        print(f"cannot {args.cmd}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
