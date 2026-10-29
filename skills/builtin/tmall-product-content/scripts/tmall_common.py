"""Shared helpers for the Tmall product-content scripts."""
from __future__ import annotations
import hashlib, json, os, sys, tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "shared" / "scripts"))
import output_paths  # noqa: E402

SECTION = "天猫商品内容"

def product_dir(brand: str, product: str, *, create: bool = False) -> Path:
    return output_paths.validate_project_dir(Path("outputs") / brand / SECTION / product, create=create)

def sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

def write_json(path, data) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        os.replace(tmp, path)
    except BaseException:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise

def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))

def open_sheet(xlsx, sheet: str):
    """One worksheet of a workbook; a missing sheet names the sheets that do exist. openpyxl is imported here (not declared in pyproject.toml)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"{xlsx}: no sheet {sheet!r}; sheets found: {wb.sheetnames}")
    return wb[sheet]

def column_index(header: list, wanted, where: str) -> list[int]:
    """Positions of the wanted header texts; a missing one names the file, the sheet and the columns actually found."""
    missing = [w for w in wanted if w not in header]
    if missing:
        raise ValueError(f"{where}: columns {missing} not found; columns found: {header}")
    return [header.index(w) for w in wanted]
