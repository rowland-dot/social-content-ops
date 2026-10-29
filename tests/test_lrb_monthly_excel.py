"""One route builds every LRB monthly plan and monthly report Excel (issue 57).

Seams (the issue's Agent Brief key interfaces and acceptance criteria): the command line of
skills/builtin/lrb-monthly-excel/scripts/lrb_excel.py, with its two subcommands `build` (the team's
draft in, the finished client file out) and `check` (the checks that fail a run, on any finished file),
plus the skill's written text. Nothing else is imported from the script.

Three kinds of input stand in for the outside world:
  - a stand-in LRB folder (`standin_lrb`), written by the tests, so the build and check logic run on
    CI, where the owner's own LRB folder does not exist. Its `verify` copies the real one's shape; its
    `style_plan` / `style_report` copy what they are given to `<out>.arrived.xlsx` (the draft exactly as
    the formatting step received it) and write a finished-looking file.
  - workbooks the tests build with openpyxl, with literal expected values typed from the issue and from
    the owner's reference files, never recomputed the way the script computes them.
  - The owner's real LRB folder, for the tests marked REAL below: skipped, with the reason, where the folder
    (or LibreOffice) is not on the machine. Client files are never copied into the repo.
"""
from __future__ import annotations

import datetime
import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

import web.app as web

ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = ROOT / "skills" / "builtin" / "lrb-monthly-excel"
SCRIPT = SKILL_DIR / "scripts" / "lrb_excel.py"

REAL_LRB = next(
    (Path(p) for p in (
        os.environ.get("LRB_DIR", ""),
        *sorted(str(p) for p in Path("/mnt/c/Users").glob("*/Claude Cowork/LRB Plan & Report")),
    ) if p and (Path(p) / "playbooks" / "scripts" / "planstyle.py").is_file()),
    None,
)
HAS_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))
REAL_ACCOUNTS = [a.strip() for a in os.environ.get("LRB_REAL_ACCOUNTS", "").split(";") if a.strip()]
# the names, as they appear in file names in the owner's LRB folder, come from this setting and stay out of the repo;
# the first one is the account whose reference files the real-folder build tests compare against
REAL_ACCOUNT = REAL_ACCOUNTS[0] if REAL_ACCOUNTS else ""
REAL_REASON = (
    "The owner's LRB folder, or LibreOffice for the formula-value step, is not on this machine; "
    "the real-folder tests run in WSL on his PC (client files are never copied into the repo)"
)
real = pytest.mark.skipif(REAL_LRB is None or not HAS_SOFFICE or os.name == "nt" or not REAL_ACCOUNT,
                          reason=REAL_REASON + "; or LRB_REAL_ACCOUNTS (account names, separated by ;) is not set")

TITLE = "Harbourside Mill 小红书11月内容规划\nHarbourside Mill content planning for November on LRB"
HEADERS = [
    "发布时间\nPublishing Time", "笔记类型\nNote Type", "笔记选题\nNote Topic", "主要内容\nMain Content",
    "预计阅读量\nEstimated Views", "预计增粉\nEstimated New Followers",
    "预计赞藏评\nEstimated Likes/Favorites/Comments",
]
TYPE = "产品推荐\nProduct Recommendation"
BODY = ("Responsible for scheduled publishing and comment management for 8 notes in November.\n"
        "Daily monitoring of DMs and comments to answer product questions.")

# literal examples from the issue's Evidence: the team draft's own cells (The owner's in/ copy, 2026-10-16)
TOPIC_PIPE_3 = ("小众茶树柠檬桃护手乳|手不干的别碰！\n"
                "Hidden-gem tea tree–lemon–peach hand cream｜Skip it if your hands don't get dry!")
TOPIC_PIPE_8 = ("木质玫瑰香|承包我一整个秋冬的护手霜！\n"
                "Woody rose scent｜The one hand cream getting me through fall & winter!")
TOPIC_PIPE_10 = ("高级清冷白花护手霜|好闻滋润全占了！\n"
                 "A refined, crisp white-floral hand cream｜Smells amazing & seriously hydrating—it does it all!")
# the triage note's own example of a line the formatting script merges
TRIAGE_EXAMPLE = "护手乳推荐\nHand lotion｜Tea Tree"
# the checks a plan runs; "fills" joined them on the owner's ruling of 2026-10-17 ("No colours, per template")
PLAN_CHECKS = {"verify", "logo", "spacer", "bilingual", "chars", "font", "fills"}


# ---- stand-in LRB folder ---------------------------------------------------------------------

STANDIN_LRBSTYLE = '''
def verify(path):
    """Same shape as the real one: fail loudly if a saved file would open anywhere but the top."""
    import openpyxl
    ws = openpyxl.load_workbook(path).worksheets[0]
    v = ws.sheet_view
    problems = []
    if (v.topLeftCell or "A1") != "A1":
        problems.append(f"opens scrolled to {v.topLeftCell}")
    if ws.freeze_panes != "A3":
        problems.append(f"freeze is {ws.freeze_panes}, expected A3")
    if v.zoomScale != 100:
        problems.append(f"zoom {v.zoomScale}")
    if problems:
        raise SystemExit(f"{path}: " + "; ".join(problems))
'''

STANDIN_STYLE = '''
import json, os, shutil
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Font
from openpyxl.worksheet.cell_range import CellRange


def _finish(src, dst, sheet_name, logo_from, kind, shift):
    if os.environ.get("STANDIN_RAISE"):
        raise RuntimeError(os.environ["STANDIN_RAISE"])
    shutil.copy(src, dst + ".arrived.xlsx")
    with open(dst + ".call.json", "w", encoding="utf-8") as f:
        json.dump({"kind": kind, "src": src, "dst": dst, "sheet": sheet_name, "logo": logo_from}, f)
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    if sheet_name:
        ws.title = sheet_name
    if shift:
        merges = [str(m) for m in ws.merged_cells.ranges]
        for m in merges:
            ws.unmerge_cells(m)
        ws.move_range(f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}", cols=1)
        for m in merges:
            cr = CellRange(m)
            cr.shift(col_shift=1)
            ws.merge_cells(cr.coord)
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10)
    ws.sheet_view.zoomScale = 100
    ws.sheet_view.topLeftCell = "A1"
    ws.freeze_panes = "A3"
    title = next(m for m in ws.merged_cells.ranges if m.min_row == 1)
    for sel in ws.sheet_view.selection:
        sel.activeCell = title.start_cell.coordinate
        sel.sqref = title.coord
    if logo_from:
        im = Image(logo_from)
        s = min(1, 84 / im.height, 280 / im.width)
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=1, row=0, colOff=0, rowOff=0),
                                  ext=XDRPositiveSize2D(int(im.width * s * 9525), int(im.height * s * 9525)))
        ws.add_image(im)
    if os.environ.get("STANDIN_FILL"):  # a formatting step that leaves a colour behind
        from openpyxl.styles import PatternFill
        ws["B3"].fill = PatternFill("solid", fgColor="FFF3EEF8")
    wb.save(dst)
    if os.environ.get("STANDIN_RAISE_AFTER_SAVE"):  # like the real one: verify() fails after the file is saved
        raise SystemExit(dst + ": " + os.environ["STANDIN_RAISE_AFTER_SAVE"])


def style_plan(src, dst, sheet_name=None, logo_from=None):
    _finish(src, dst, sheet_name, logo_from, "plan", True)


def style_report(src, dst, sheet_name=None, logo=None):
    _finish(src, dst, sheet_name, logo, "report", False)
'''


def standin_lrb(tmp_path: Path, name: str = "lrb") -> Path:
    scripts = tmp_path / name / "playbooks" / "scripts"
    scripts.mkdir(parents=True)
    (scripts / "lrbstyle.py").write_text(STANDIN_LRBSTYLE, encoding="utf-8")
    (scripts / "planstyle.py").write_text(STANDIN_STYLE, encoding="utf-8")
    (scripts / "reportstyle.py").write_text(STANDIN_STYLE, encoding="utf-8")
    return tmp_path / name


# ---- workbooks the tests build -----------------------------------------------------------------


def png(path: Path, width: int, height: int) -> Path:
    from PIL import Image

    Image.new("RGBA", (width, height), (106, 75, 153, 255)).save(path)
    return path


def finished_plan(path: Path, *, logo=(280, 60, 600, 128), spacers=1, mutate=None) -> Path:
    """A plan in the finished shape (gutter column A, table B:H, one spacer row, work plan block)."""
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nov 2026"
    arial = Font(name="Arial", size=10)

    def put(coord, value):
        ws[coord] = value
        ws[coord].font = arial
        ws[coord].alignment = Alignment(wrap_text=True, vertical="center")

    put("B1", TITLE)
    ws.merge_cells("B1:H1")
    for col, text in zip("BCDEFGH", HEADERS):
        put(f"{col}2", text)
    for i in range(8):
        r = 3 + i
        put(f"B{r}", datetime.datetime(2026, 11, 1 + i))
        put(f"C{r}", TYPE)
        put(f"D{r}", f"小众护手乳第{i + 1}条\nHidden-gem hand cream note {i + 1}")
        put(f"E{r}", "茶树护手乳\nTea Tree Hand Lotion")
        put(f"F{r}", 1000)
        put(f"G{r}", 20)
        put(f"H{r}", 20)
    wp = 11 + spacers
    put(f"B{wp}", "Work plan breakdown")
    ws.merge_cells(f"B{wp}:B{wp + 1}")
    for k, label in enumerate(("Daily Routine", "Content Testing")):
        put(f"C{wp + k}", label)
        put(f"D{wp + k}", BODY)
        ws.merge_cells(f"D{wp + k}:H{wp + k}")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_view.topLeftCell = "A1"
    ws.freeze_panes = "A3"
    for sel in ws.sheet_view.selection:
        sel.activeCell = "B1"
        sel.sqref = "B1:H1"
    if logo is not None:
        w, h, native_w, native_h = logo
        im = Image(png(path.with_suffix(".logo.png"), native_w, native_h))
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=1, row=0, colOff=6 * 9525, rowOff=0),
                                  ext=XDRPositiveSize2D(int(w * 9525), int(h * 9525)))
        ws.add_image(im)
    if mutate:
        mutate(ws)
    wb.save(path)
    return path


def draft_plan(path: Path, *, start_col=1, dup_heading=False, blank_rows=1, topic_rows=None, mutate=None) -> Path:
    """A plan the way the team sends it: no gutter column, Calibri/Song fonts, draft shape.

    start_col=2 is the DingTalk copy's shape (column A empty, table from column B).
    """
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter as L

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Harbourside Mill 11月内容规划"
    s = start_col

    def put(r, c, value):
        cell = ws.cell(r, c, value)
        cell.font = Font(name="宋体", size=10)
        return cell

    put(1, s, TITLE)
    ws.merge_cells(f"{L(s)}1:{L(s + 6)}1")
    for k, text in enumerate(HEADERS):
        put(2, s + k, text)
    topic_rows = topic_rows or {}
    for i in range(8):
        r = 3 + i
        put(r, s, f"{i + 1:02d}/11/2026")
        put(r, s + 1, TYPE)
        put(r, s + 2, topic_rows.get(r, f"小众护手乳第{i + 1}条\nHidden-gem hand cream note {i + 1}"))
        put(r, s + 3, "茶树护手乳\nTea Tree Hand Lotion")
        put(r, s + 4, "1000")
        put(r, s + 5, "20")
        put(r, s + 6, "20")
    wp = 11 + blank_rows
    if dup_heading:
        put(wp, s, "工作计划分解 Work Plan Breakdown")
        ws.merge_cells(f"{L(s)}{wp}:{L(s + 6)}{wp}")
        wp += 1
    put(wp, s, "Work plan breakdown")
    ws.merge_cells(f"{L(s)}{wp}:{L(s)}{wp + 1}")
    for k, label in enumerate(("Daily Routine", "Content Testing")):
        put(wp + k, s + 1, label)
        put(wp + k, s + 2, BODY)
        ws.merge_cells(f"{L(s + 2)}{wp + k}:{L(s + 6)}{wp + k}")
        ws.row_dimensions[wp + k].height = 55
    if mutate:
        mutate(ws)
    wb.save(path)
    return path


# ---- running the script ------------------------------------------------------------------------


class Run:
    def __init__(self, proc):
        self.code = proc.returncode
        self.out = proc.stdout
        self.err = proc.stderr
        self.lines = proc.stdout.splitlines()

    def tagged(self, tag):
        return [ln for ln in self.lines if ln.split(" ", 1)[0] == tag]

    @property
    def failed(self) -> set[str]:
        return {ln.split()[1].rstrip(":") for ln in self.lines if ln.startswith("FAIL ")}

    @property
    def passed(self) -> set[str]:
        return {ln.split()[1] for ln in self.lines if ln.startswith("PASS ")}


def run(*args, env_extra=None, project=None) -> Run:
    env = {k: v for k, v in os.environ.items() if k not in ("PYTHONDONTWRITEBYTECODE", "LRB_DIR")}
    env["PYTHONIOENCODING"] = "utf-8"
    if project is not None:
        env["EASEL_ROOT"] = str(project)
    env.update(env_extra or {})
    proc = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True,
                          text=True, encoding="utf-8", timeout=900, env=env)
    return Run(proc)


def check(path, lrb, kind="plan", **kw) -> Run:
    return run("check", "--kind", kind, str(path), "--lrb-dir", str(lrb), **kw)


@pytest.fixture
def lrb(tmp_path):
    return standin_lrb(tmp_path)


@pytest.fixture
def project(tmp_path):
    topic = tmp_path / "project" / "outputs" / "Harbourside Mill" / "2026-11 月度计划"
    topic.mkdir(parents=True)
    return tmp_path / "project"


def topic_dir(project: Path) -> Path:
    return project / "outputs" / "Harbourside Mill" / "2026-11 月度计划"


def build(source, out, lrb, project, *extra, kind="plan", sheet="Nov 2026", logo=None, env_extra=None) -> Run:
    args = ["build", "--kind", kind, "--source", str(source), "--out", str(out), "--sheet", sheet,
            "--lrb-dir", str(lrb)]
    if logo:
        args += ["--logo", str(logo)]
    return run(*args, *extra, project=project, env_extra=env_extra)


def cells(path, sheet=0):
    import openpyxl

    ws = openpyxl.load_workbook(path).worksheets[sheet]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ==== check: the checks that fail a run ========================================================


def test_check_passes_a_correctly_formatted_plan(tmp_path, lrb):
    r = check(finished_plan(tmp_path / "ok.xlsx"), lrb)
    assert r.code == 0, r.out + r.err
    assert r.passed == PLAN_CHECKS
    assert not r.failed
    assert r.lines[-1] == "RESULT OK"


def _scroll(ws):
    ws.sheet_view.topLeftCell = "A20"


def _set(coord, value):
    def mutate(ws):
        ws[coord] = value

    return mutate


def _fill(coord, argb):
    def mutate(ws):
        from openpyxl.styles import PatternFill

        ws[coord].fill = PatternFill("solid", fgColor=argb)

    return mutate


def _calibri(ws):
    from openpyxl.styles import Font

    ws["E4"].font = Font(name="Calibri", size=10)


BREAKS = {
    # id: (name of the one check that must fail, finished_plan keyword arguments)
    "logo-128-px-tall": ("logo", {"logo": (600, 128, 600, 128)}),
    "logo-too-wide": ("logo", {"logo": (300, 64, 600, 128)}),
    "logo-distorted": ("logo", {"logo": (280, 84, 600, 128)}),
    "no-logo": ("logo", {"logo": None}),
    "two-spacer-rows": ("spacer", {"spacers": 2}),
    "no-spacer-row": ("spacer", {"spacers": 0}),
    "merged-bilingual-cell": ("bilingual", {"mutate": _set("D3", TRIAGE_EXAMPLE.replace("\n", ""))}),
    "english-above-chinese": ("bilingual", {"mutate": _set("D4", "Hand lotion for dry hands\n护手乳推荐")}),
    "stray-figure-hyphen": ("chars", {"mutate": _set("D5", "护手乳\nanti\u2011aging hand lotion")}),
    "non-breaking-space": ("chars", {"mutate": _set("D6", "护手乳\nhand\u00a0lotion")}),
    "joined-emoji": ("chars", {"mutate": _set("D7", "护手乳\nrunner \U0001F3C3\u200d\u2642")}),
    "non-arial-font": ("font", {"mutate": _calibri}),
    "scrolled-view": ("verify", {"mutate": _scroll}),
    "coloured-fill": ("fills", {"mutate": _fill("B12", "FFF3EEF8")}),
    "white-fill": ("fills", {"mutate": _fill("D4", "FFFFFFFF")}),
}


@pytest.mark.parametrize("name,expected,kwargs", [(k, v[0], v[1]) for k, v in BREAKS.items()],
                         ids=list(BREAKS))
def test_check_fails_a_file_made_to_break_exactly_one_check(tmp_path, lrb, name, expected, kwargs):
    r = check(finished_plan(tmp_path / f"{name}.xlsx", **kwargs), lrb)
    assert r.code == 1, r.out + r.err
    assert r.failed == {expected}, r.out
    assert r.passed == PLAN_CHECKS - {expected}
    assert r.lines[-1] == "RESULT FAILED 1"


def test_check_names_the_cell_and_the_reason(tmp_path, lrb):
    r = check(finished_plan(tmp_path / "x.xlsx", mutate=_set("D3", TRIAGE_EXAMPLE.replace("\n", ""))), lrb)
    line = next(ln for ln in r.lines if ln.startswith("FAIL bilingual"))
    assert "D3" in line
    r = check(finished_plan(tmp_path / "y.xlsx", logo=(600, 128, 600, 128)), lrb)
    line = next(ln for ln in r.lines if ln.startswith("FAIL logo"))
    assert "128" in line and "84" in line


def test_check_does_not_mistake_a_work_plan_paragraph_with_a_chinese_word_for_a_merged_cell(tmp_path, lrb):
    # approved plans carry English paragraphs in the work plan block that quote a Chinese word
    body = "Ride seasonal search demand rather than sales events on 天猫 and the heating season in autumn."
    path = finished_plan(tmp_path / "w.xlsx", mutate=_set("D12", body))
    r = check(path, lrb)
    assert r.code == 0 and "bilingual" in r.passed, r.out


def test_check_names_every_colour_and_the_cells_it_sits_on_when_a_file_has_fills(tmp_path, lrb):
    from openpyxl.styles import PatternFill

    def fills(ws):
        ws["B12"].fill = PatternFill("solid", fgColor="FFF3EEF8")
        for col in "BCDEFGH":
            ws[f"{col}3"].fill = PatternFill("solid", fgColor="FFFFFFFF")

    r = check(finished_plan(tmp_path / "f.xlsx", mutate=fills), lrb)
    assert r.code == 1 and r.failed == {"fills"}, r.out
    line = next(ln for ln in r.lines if ln.startswith("FAIL fills"))
    assert "F3EEF8" in line and "B12" in line, line
    assert "FFFFFF 7 cells in the box B3:H3" in line, line  # a long list is given as the box that holds it
    assert r.lines[-1] == "RESULT FAILED 1"


def test_check_counts_a_fill_on_an_empty_styled_cell_and_a_theme_colour_too(tmp_path, lrb):
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color

    def fills(ws):
        ws["J20"].fill = PatternFill("solid", fgColor="FFABCDEF")  # no value in the cell
        ws["I5"].fill = PatternFill("solid", fgColor=Color(theme=0))  # a theme colour, not an rgb one

    r = check(finished_plan(tmp_path / "e.xlsx", mutate=fills), lrb)
    line = next(ln for ln in r.lines if ln.startswith("FAIL fills"))
    assert "ABCDEF" in line and "J20" in line and "theme0" in line and "I5" in line, r.out


def test_check_passes_fills_when_there_are_none(tmp_path, lrb):
    r = check(finished_plan(tmp_path / "n.xlsx"), lrb)
    assert "fills" in r.passed and not r.tagged("FILLS"), r.out


def test_spacer_rule_is_for_plans_only(tmp_path, lrb):
    path = finished_plan(tmp_path / "two.xlsx", spacers=2)
    assert check(path, lrb, kind="plan").failed == {"spacer"}
    r = check(path, lrb, kind="report")
    assert "spacer" not in r.passed | r.failed and r.code == 0, r.out


@pytest.mark.parametrize("why,args,needle", [
    ("the file does not exist", lambda t, lrb: ["check", "--kind", "plan", str(t / "nope.xlsx"), "--lrb-dir", str(lrb)],
     "cannot check"),
    ("the LRB folder does not exist", lambda t, lrb: ["check", "--kind", "plan", str(t / "ok.xlsx"), "--lrb-dir", str(t / "gone")],
     "cannot check"),
    ("the kind is unknown", lambda t, lrb: ["check", "--kind", "deck", str(t / "ok.xlsx"), "--lrb-dir", str(lrb)],
     "invalid choice"),
])
def test_a_check_that_cannot_run_is_exit_2_and_says_why(tmp_path, lrb, why, args, needle):
    finished_plan(tmp_path / "ok.xlsx")
    r = run(*args(tmp_path, lrb))
    assert r.code == 2, (why, r.out, r.err)
    assert needle in r.err and "RESULT" not in r.out, (why, r.err)


def test_a_lrb_folder_without_verify_is_not_a_pass(tmp_path):
    finished_plan(tmp_path / "ok.xlsx")
    bare = tmp_path / "bare"
    (bare / "playbooks" / "scripts").mkdir(parents=True)
    (bare / "playbooks" / "scripts" / "lrbstyle.py").write_text("X = 1\n", encoding="utf-8")
    r = check(tmp_path / "ok.xlsx", bare)
    assert r.code == 2 and "cannot check" in r.err and "verify" in r.err, r.err


def test_check_leaves_the_file_and_the_lrb_folder_untouched(tmp_path, lrb):
    path = finished_plan(tmp_path / "ok.xlsx")
    before = sha(path)
    files_before = sorted(p.name for p in (lrb / "playbooks" / "scripts").iterdir())
    assert check(path, lrb).code == 0
    assert sha(path) == before
    assert sorted(p.name for p in (lrb / "playbooks" / "scripts").iterdir()) == files_before, \
        "importing the LRB scripts must not write bytecode into the owner's folder"


# ==== build: the team's draft in, the finished client file out ===================================


def test_build_hands_the_formatting_step_a_copy_never_the_original(tmp_path, lrb, project):
    draft = draft_plan(tmp_path / "draft.xlsx")
    before = sha(draft)
    logo = png(tmp_path / "logo.png", 600, 128)
    out = topic_dir(project) / "Harbourside Mill LRB Plan Nov 2026.xlsx"
    r = build(draft, out, lrb, project, logo=logo)
    assert r.code == 0, r.out + r.err
    assert sha(draft) == before
    call = json.loads(Path(str(out) + ".call.json").read_text(encoding="utf-8"))
    assert call["kind"] == "plan" and call["sheet"] == "Nov 2026" and call["logo"] == str(logo)
    assert call["src"] != str(draft) and call["dst"] == str(out)
    assert r.lines[-1] == "RESULT OK"
    assert r.passed == PLAN_CHECKS


def test_build_moves_a_table_that_starts_in_column_b_to_column_a_first(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", start_col=2), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert any("column B to column A" in ln for ln in r.tagged("SHAPE")), r.out
    arrived = cells(Path(str(out) + ".arrived.xlsx"))
    assert arrived["A1"] == TITLE and arrived["A2"] == HEADERS[0]
    assert not [c for c in arrived if c.startswith("H")], "nothing is left in a column past G"
    done = cells(out)
    assert done["B1"] == TITLE and done["D3"].startswith("小众护手乳第1条")


def test_build_does_not_move_a_table_that_already_starts_in_column_a(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0 and not r.tagged("SHAPE"), r.out
    assert cells(Path(str(out) + ".arrived.xlsx"))["A1"] == TITLE


def test_build_removes_the_duplicate_work_plan_heading_row(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", start_col=2, dup_heading=True), out, lrb, project,
              logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert any("duplicate" in ln and "row 12" in ln and "工作计划分解 Work Plan Breakdown" in ln
               for ln in r.tagged("SHAPE")), r.out
    arrived = cells(Path(str(out) + ".arrived.xlsx"))
    work_rows = [c for c, v in arrived.items() if c.startswith("A") and isinstance(v, str) and "Work" in v]
    assert work_rows == ["A12"], work_rows
    assert arrived["B12"] == "Daily Routine" and arrived["B13"] == "Content Testing"
    assert "工作计划分解 Work Plan Breakdown" not in arrived.values()
    assert "spacer" in r.passed and r.code == 0


def test_build_keeps_the_work_plan_merges_after_removing_a_row(tmp_path, lrb, project):
    import openpyxl

    out = topic_dir(project) / "p.xlsx"
    build(draft_plan(tmp_path / "d.xlsx", dup_heading=True), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    ws = openpyxl.load_workbook(str(out) + ".arrived.xlsx").worksheets[0]
    assert sorted(str(m) for m in ws.merged_cells.ranges) == sorted(["A1:G1", "A12:A13", "C12:G12", "C13:G13"])
    assert ws.max_row == 13
    # the draft gave its last work-plan row (14 before the removal) a height; it must not stay on the empty row
    assert ws.row_dimensions[13].height == 55 and ws.row_dimensions[14].height is None


def test_build_collapses_extra_blank_rows_between_the_schedule_and_the_work_plan(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", blank_rows=3), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert any("2 extra blank row" in ln for ln in r.tagged("SHAPE")), r.out
    arrived = cells(Path(str(out) + ".arrived.xlsx"))
    assert arrived["A12"] == "Work plan breakdown"
    assert "spacer" in r.passed


def test_build_keeps_a_full_width_english_line_as_line_two_and_offers_the_colon(tmp_path, lrb, project):
    rows = {3: TOPIC_PIPE_3, 8: TOPIC_PIPE_8, 10: TOPIC_PIPE_10}
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", topic_rows=rows), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    # the formatting step never sees the full-width character, which it would count as Chinese
    arrived = cells(Path(str(out) + ".arrived.xlsx"))
    assert "｜" not in arrived["C3"] and "\n" in arrived["C3"]
    # ...and the finished file carries the draft's exact text, character for character
    done = cells(out)
    assert done["D3"] == TOPIC_PIPE_3 and done["D8"] == TOPIC_PIPE_8 and done["D10"] == TOPIC_PIPE_10
    assert [ln.split()[1].rstrip(":") for ln in r.tagged("KEPT-LINE")] == ["D3", "D8", "D10"]
    # only the English line of each cell is kept and offered; the Chinese line above it (which also ends in a
    # full-width "！") is neither
    assert r.tagged("KEPT-LINE")[0].startswith("KEPT-LINE D3: Hidden-gem tea tree")
    assert len(r.tagged("OFFERED")) == 3, r.tagged("OFFERED")
    offered = {ln.split()[1].rstrip(":"): ln for ln in r.tagged("OFFERED")}
    assert set(offered) == {"D3", "D8", "D10"}
    assert "hand cream: Skip it if your hands don't get dry!" in offered["D3"]
    assert "｜" in offered["D3"].split("=>")[0] and "｜" not in offered["D3"].split("=>")[1]
    assert not r.tagged("APPLIED"), "offering a content change is not applying it"


def test_build_cell_with_only_ordinary_characters_is_not_reported_as_kept(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert not r.tagged("KEPT-LINE") and not r.tagged("OFFERED")


def test_build_keeps_a_line_but_offers_nothing_when_there_is_no_half_width_form_to_suggest(tmp_path, lrb, project):
    cell = "买一送一\nBuy one 【now】 get one free"  # 【 】 are full-width brackets with no plain-ASCII twin
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", topic_rows={4: cell}), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert [ln.split()[1].rstrip(":") for ln in r.tagged("KEPT-LINE")] == ["D4"]
    assert not r.tagged("OFFERED")
    assert cells(out)["D4"] == cell


def test_build_triage_example_cell_keeps_its_two_lines(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    build(draft_plan(tmp_path / "d.xlsx", topic_rows={5: TRIAGE_EXAMPLE}), out, lrb, project,
          logo=png(tmp_path / "l.png", 600, 128))
    assert cells(out)["D5"] == TRIAGE_EXAMPLE


def test_build_applies_only_the_listed_content_changes_and_lists_them(tmp_path, lrb, project):
    draft = draft_plan(tmp_path / "d.xlsx")
    changes = tmp_path / "changes.json"
    ruling = 'the owner, 2026-10-17: "use Australian spelling"'
    out = topic_dir(project) / "p.xlsx"
    changes.write_text(json.dumps([{
        "old": "小众护手乳第4条\nHidden-gem hand cream note 4",
        "new": "小众护手乳第4条\nHidden-gem moisturising hand cream note 4", "ruling": ruling}],
        ensure_ascii=False), encoding="utf-8")
    r = build(draft, out, lrb, project, "--changes", changes, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    done, src = cells(out), cells(draft)
    assert done["D6"] == "小众护手乳第4条\nHidden-gem moisturising hand cream note 4"
    changed = {c for c in done if c != "D6" and done[c] != _from_source(src, c)}
    assert not changed, f"cells changed that were not listed: {sorted(changed)}"
    applied = r.tagged("APPLIED")
    assert len(applied) == 1 and "D6" in applied[0] and ruling in applied[0]
    assert "Hidden-gem hand cream note 4" in applied[0] and "moisturising" in applied[0]


def _from_source(src, out_coord):
    """The draft's cell behind a finished-file coordinate (the finished file adds a gutter column)."""
    from openpyxl.utils import column_index_from_string, get_column_letter

    letters = "".join(ch for ch in out_coord if ch.isalpha())
    digits = "".join(ch for ch in out_coord if ch.isdigit())
    return src.get(f"{get_column_letter(column_index_from_string(letters) - 1)}{digits}")


def test_build_writes_the_content_changes_to_a_separate_file_when_asked(tmp_path, lrb, project):
    changes = tmp_path / "changes.json"
    changes.write_text(json.dumps([{
        "old": "小众护手乳第4条\nHidden-gem hand cream note 4", "new": "小众护手乳第4条\nHidden-gem note 4",
        "ruling": 'the owner: "shorter"'}], ensure_ascii=False), encoding="utf-8")
    listing = topic_dir(project) / "p content changes.md"
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, "--changes", changes, "--changes-out", listing,
              logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    text = listing.read_text(encoding="utf-8")
    assert "D6" in text and "Hidden-gem hand cream note 4" in text and "Hidden-gem note 4" in text
    assert 'the owner: "shorter"' in text


@pytest.mark.parametrize("why,old", [
    ("the old text is not in the draft any more", "this text is nowhere in the draft"),
    ("the old text is in more than one cell", "茶树护手乳\nTea Tree Hand Lotion"),
])
def test_build_refuses_a_content_change_it_cannot_place_exactly(tmp_path, lrb, project, why, old):
    changes = tmp_path / "changes.json"
    changes.write_text(json.dumps([{"old": old, "new": "x", "ruling": 'the owner: "yes"'}], ensure_ascii=False),
                       encoding="utf-8")
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, "--changes", changes, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 2, (why, r.out, r.err)
    assert "content change 1" in r.err and "matches" in r.err and not out.exists(), (why, r.err)


def test_build_refuses_a_content_change_without_a_ruling(tmp_path, lrb, project):
    changes = tmp_path / "changes.json"
    changes.write_text(json.dumps([{"old": "a", "new": "b", "ruling": ""}]), encoding="utf-8")
    r = build(draft_plan(tmp_path / "d.xlsx"), topic_dir(project) / "p.xlsx", lrb, project, "--changes", changes)
    assert r.code == 2 and "ruling" in r.err


# ---- round 3: a change can name its cell; a ruling is the owner's own quoted words -------------------------

NOTE_27 = "礼物推荐\nGift Recommendation"  # the owner's "Easel's 27 Nov" ruling: 产品推荐 / Product Recommendation -> this
HERS = 'the owner, 2026-10-17: "Hers fully, Easel\'s 27 Nov (Recommended)"'


def change_file(tmp_path: Path, *changes: dict) -> Path:
    path = tmp_path / "changes.json"
    path.write_text(json.dumps(list(changes), ensure_ascii=False), encoding="utf-8")
    return path


def build_with(tmp_path, lrb, project, *changes, draft=None, kind="plan", out_name="p.xlsx", extra=()):
    out = topic_dir(project) / out_name
    r = build(draft or draft_plan(tmp_path / "d.xlsx"), out, lrb, project, "--changes", change_file(tmp_path, *changes),
              *extra, kind=kind, logo=png(tmp_path / "l.png", 600, 128))
    return r, out


def test_build_applies_a_change_to_the_cell_it_names_when_the_same_text_sits_in_other_cells(tmp_path, lrb, project):
    # the text is in all eight schedule rows; the change names the 27 Nov row's cell in the finished file
    r, out = build_with(tmp_path, lrb, project, {"cell": "C10", "old": TYPE, "new": NOTE_27, "ruling": HERS})
    assert r.code == 0, r.out + r.err
    done = cells(out)
    assert done["C10"] == NOTE_27
    assert [c for c in done if c.startswith("C") and done[c] == TYPE] == [f"C{n}" for n in range(3, 10)]
    applied = r.tagged("APPLIED")
    assert len(applied) == 1 and applied[0].startswith("APPLIED C10:") and HERS in applied[0], r.out


def test_build_finds_the_cell_by_row_date_and_column(tmp_path, lrb, project):
    for date in ("2026-11-08", "08/11/2026"):  # the test draft's last row is 08/11/2026; ISO is accepted too
        r, out = build_with(tmp_path, lrb, project, {"date": date, "column": "C", "old": TYPE, "new": NOTE_27,
                                                     "ruling": HERS}, out_name=f"{date[-2:]}.xlsx")
        assert r.code == 0, r.out + r.err
        assert cells(out)["C10"] == NOTE_27 and cells(out)["C9"] == TYPE
        assert r.tagged("APPLIED")[0].startswith("APPLIED C10:")


def test_build_a_change_that_names_a_cell_works_on_a_column_b_draft_and_on_a_report(tmp_path, lrb, project):
    r, out = build_with(tmp_path, lrb, project, {"cell": "C10", "old": TYPE, "new": NOTE_27, "ruling": HERS},
                        draft=draft_plan(tmp_path / "b.xlsx", start_col=2))
    assert r.code == 0 and cells(out)["C10"] == NOTE_27, r.out + r.err
    r, out = build_with(tmp_path, lrb, project, {"cell": "B10", "old": TYPE, "new": NOTE_27, "ruling": HERS},
                        kind="report", out_name="r.xlsx")  # a report keeps the draft's own coordinates
    assert r.code == 0 and cells(out)["B10"] == NOTE_27 and cells(out)["B9"] == TYPE, r.out + r.err


@pytest.mark.parametrize("why,entry,needle", [
    ("the named cell holds other text", {"cell": "C9", "old": "something else", "new": "x"}, "C9"),
    ("the named cell is in the gutter column", {"cell": "A10", "old": TYPE, "new": "x"}, "A10"),
    ("no row has that date", {"date": "2026-11-28", "column": "C", "old": TYPE, "new": "x"}, "2026-11-28"),
    ("a date with no column", {"date": "2026-11-27", "old": TYPE, "new": "x"}, "needs a \"column\""),
    ("a date with the gutter column", {"date": "2026-11-08", "column": "A", "old": TYPE, "new": "x"}, "gutter"),
    ("a date that is not a date", {"date": "next week", "column": "C", "old": TYPE, "new": "x"}, "'next week' is not a date"),
    ("a cell and a date together", {"cell": "C10", "date": "2026-11-27", "column": "C", "old": TYPE, "new": "x"}, "not both"),
    ("the address is not an address", {"cell": "ten", "old": TYPE, "new": "x"}, "'ten' is not a cell address"),
])
def test_build_refuses_a_change_whose_address_does_not_lead_to_its_old_text(tmp_path, lrb, project, why, entry, needle):
    r, out = build_with(tmp_path, lrb, project, {**entry, "ruling": HERS})
    assert r.code == 2, (why, r.out, r.err)
    assert "content change 1" in r.err and needle in r.err and not out.exists(), (why, r.err)
    if entry["old"] == "something else":
        assert "Product Recommendation" in r.err, "the message shows what the cell actually holds"


def test_build_a_date_that_names_two_rows_is_refused_rather_than_guessed(tmp_path, lrb, project):
    draft = draft_plan(tmp_path / "d.xlsx", mutate=lambda ws: ws.__setitem__("A4", "01/11/2026"))  # rows 3 and 4 share a date
    r, out = build_with(tmp_path, lrb, project, {"date": "2026-11-01", "column": "C", "old": TYPE, "new": NOTE_27,
                                                 "ruling": HERS}, draft=draft)
    assert r.code == 2 and not out.exists(), r.out
    assert "2 schedule rows have the date 2026-11-01" in r.err, r.err


def test_build_ambiguous_text_without_an_address_names_every_matching_cell(tmp_path, lrb, project):
    r, out = build_with(tmp_path, lrb, project, {"old": TYPE, "new": NOTE_27, "ruling": HERS})
    assert r.code == 2 and not out.exists(), r.out
    assert "matches 8 cells" in r.err
    assert all(f"C{n}" in r.err for n in range(3, 11)), r.err  # the finished file's cell names, ready to copy
    assert '"cell"' in r.err, "the message says how to name one"


def test_build_unique_text_still_needs_no_address(tmp_path, lrb, project):
    unique = "小众护手乳第2条\nHidden-gem hand cream note 2"
    r, out = build_with(tmp_path, lrb, project, {"old": unique, "new": "x\ny", "ruling": HERS})
    assert r.code == 0 and cells(out)["D4"] == "x\ny", r.out + r.err


@pytest.mark.parametrize("ruling", [
    "The owner approved it",                    # a summary, no quoted words
    "The owner: \"\"",                          # quotation marks around nothing
    "The owner: \" \"",
])
def test_build_refuses_a_ruling_that_is_not_ros_quoted_words(tmp_path, lrb, project, ruling):
    r, out = build_with(tmp_path, lrb, project, {"cell": "C10", "old": TYPE, "new": NOTE_27, "ruling": ruling})
    assert r.code == 2 and not out.exists(), r.out
    assert "content change 1" in r.err and "quot" in r.err, r.err


def test_build_carries_ros_quoted_words_unchanged_into_the_applied_line_and_the_list(tmp_path, lrb, project):
    listing = topic_dir(project) / "list.md"
    chinese = "负责人，2026-10-17：「护手乳那条改成礼物推荐」"
    r, out = build_with(tmp_path, lrb, project,
                        {"cell": "C10", "old": TYPE, "new": NOTE_27, "ruling": HERS},
                        {"cell": "C5", "old": TYPE, "new": NOTE_27, "ruling": chinese},
                        extra=("--changes-out", listing))
    assert r.code == 0, r.out + r.err
    text = listing.read_text(encoding="utf-8")
    for words in ("Hers fully, Easel's 27 Nov (Recommended)", "护手乳那条改成礼物推荐"):
        assert words in text and any(words in ln for ln in r.tagged("APPLIED")), words
    assert "| C10 |" in text and "| C5 |" in text


def test_build_only_writes_inside_outputs_topic_folders(tmp_path, lrb, project):
    for bad in (tmp_path / "loose.xlsx", project / "outputs" / "loose.xlsx"):
        r = build(draft_plan(tmp_path / "d.xlsx"), bad, lrb, project)
        assert r.code == 2 and "cannot build" in r.err and not bad.exists(), (bad, r.err)


def fill_cells(path) -> dict[str, str]:
    """Every cell of the first sheet that carries any fill, with the fill's colour (theme colours as themeN)."""
    import openpyxl

    ws = openpyxl.load_workbook(path).worksheets[0]
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.fill is not None and c.fill.fill_type is not None:
                col = c.fill.fgColor
                out[c.coordinate] = col.rgb[-6:] if col.type == "rgb" else f"theme{col.theme}"
    return out


def _draft_fills(ws):
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color

    ws["A12"].fill = PatternFill("solid", fgColor="FFF3EEF8")  # the work-plan heading, as in the team's EA draft
    ws["B12"].fill = PatternFill("solid", fgColor="FFF3EEF8")
    for r in (3, 4, 5):
        ws[f"A{r}"].fill = PatternFill("solid", fgColor=Color(theme=0))  # a "white" that is a fill all the same
    ws["H20"].fill = PatternFill("solid", fgColor="FFF9F6FC")  # an empty cell that only carries a colour


@pytest.mark.parametrize("kind,shift", [("plan", 1), ("report", 0)])
def test_build_clears_every_fill_the_draft_carries_and_lists_what_it_cleared(tmp_path, lrb, project, kind, shift):
    draft = draft_plan(tmp_path / "d.xlsx", mutate=_draft_fills)
    before = sha(draft)
    assert fill_cells(draft), "the test draft must carry fills"
    out = topic_dir(project) / "p.xlsx"
    r = build(draft, out, lrb, project, kind=kind, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert sha(draft) == before, "the team's draft is never changed"
    assert fill_cells(Path(str(out) + ".arrived.xlsx")) == {}, "the formatting step must receive a draft with no colours"
    assert fill_cells(out) == {}
    assert "fills" in r.passed and r.lines[-1] == "RESULT OK", r.out
    col = lambda letter: chr(ord(letter) + shift)  # noqa: E731 — the plan's gutter column moves every cell one right
    cleared = " ".join(r.tagged("CLEARED"))
    assert "F3EEF8" in cleared and f"{col('A')}12" in cleared and f"{col('B')}12" in cleared, r.out
    assert "theme0" in cleared and f"{col('A')}3" in cleared and f"{col('A')}5" in cleared, r.out
    assert "F9F6FC" in cleared and f"{col('H')}20" in cleared, r.out
    # the clean-up is formatting only: every value is the draft's own
    done, src = cells(out), cells(draft)
    for coord, value in src.items():
        assert done[f"{col(coord[0])}{coord[1:]}"] == value, coord


def test_build_prints_nothing_about_fills_when_the_draft_has_none(tmp_path, lrb, project):
    r = build(draft_plan(tmp_path / "d.xlsx"), topic_dir(project) / "p.xlsx", lrb, project,
              logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0 and "fills" in r.passed and not r.tagged("CLEARED"), r.out


def test_build_clears_the_fills_of_a_column_b_draft_after_moving_it(tmp_path, lrb, project):
    def fill(ws):
        from openpyxl.styles import PatternFill

        ws["B12"].fill = PatternFill("solid", fgColor="FFF3EEF8")  # the work-plan heading in a column-B draft

    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx", start_col=2, mutate=fill), out, lrb, project,
              logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0, r.out + r.err
    assert fill_cells(out) == {} and any("B12" in ln for ln in r.tagged("CLEARED")), r.out


def test_build_a_colour_the_formatting_step_leaves_behind_fails_the_run(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128),
              env_extra={"STANDIN_FILL": "1"})
    assert r.code == 1 and r.failed == {"fills"}, r.out
    assert "B3" in next(ln for ln in r.lines if ln.startswith("FAIL fills"))


def test_build_a_formatting_step_that_raises_is_a_failed_run_not_a_pass(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, env_extra={"STANDIN_RAISE": "boom from planstyle"})
    assert r.code == 1, r.out + r.err
    assert any(ln.startswith("FAIL format") and "boom from planstyle" in ln for ln in r.lines), r.out
    assert r.lines[-1].startswith("RESULT FAILED")


def test_build_a_file_that_fails_verify_after_it_was_saved_is_not_left_behind(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, env_extra={"STANDIN_RAISE_AFTER_SAVE": "opens scrolled to A20"})
    assert r.code == 1, r.out + r.err
    assert any(ln.startswith("FAIL format") and "opens scrolled to A20" in ln for ln in r.lines), r.out
    assert Path(str(out) + ".arrived.xlsx").is_file(), "the stand-in did save its file before failing"
    assert not out.exists(), "a file that did not pass verify() must not stay in the project's outputs"


def test_build_a_finished_file_that_breaks_a_check_fails_the_run(tmp_path, lrb, project):
    # no --logo and the draft has none: the formatting step cannot invent one
    r = build(draft_plan(tmp_path / "d.xlsx"), topic_dir(project) / "p.xlsx", lrb, project)
    assert r.code == 1 and r.failed == {"logo"}, r.out
    assert "--logo" in next(ln for ln in r.lines if ln.startswith("FAIL logo"))


def test_build_a_report_goes_to_the_report_script_and_is_not_reshaped(tmp_path, lrb, project):
    out = topic_dir(project) / "r.xlsx"
    # column B start would be moved if it were a plan; the report keeps its own gutter, so cells keep their coordinates
    draft = draft_plan(tmp_path / "d.xlsx", start_col=2, topic_rows={3: TOPIC_PIPE_3})
    r = build(draft, out, lrb, project, kind="report", sheet="Sep 2026", logo=png(tmp_path / "l.png", 600, 128))
    assert json.loads(Path(str(out) + ".call.json").read_text(encoding="utf-8"))["kind"] == "report"
    assert not r.tagged("SHAPE")
    assert cells(Path(str(out) + ".arrived.xlsx"))["B1"] == TITLE
    assert [ln.split()[1].rstrip(":") for ln in r.tagged("KEPT-LINE")] == ["D3"], r.out
    assert cells(out)["D3"] == TOPIC_PIPE_3


def test_build_writes_no_bytecode_into_the_lrb_folder(tmp_path, lrb, project):
    out = topic_dir(project) / "p.xlsx"
    r = build(draft_plan(tmp_path / "d.xlsx"), out, lrb, project, logo=png(tmp_path / "l.png", 600, 128))
    assert r.code == 0 and Path(str(out) + ".call.json").is_file(), "the stand-in formatting step must have run"
    assert not list(lrb.rglob("__pycache__")) and not list(lrb.rglob("*.pyc"))


def test_build_without_the_lrb_folder_cannot_build_and_says_so(tmp_path, project):
    r = build(draft_plan(tmp_path / "d.xlsx"), topic_dir(project) / "p.xlsx", tmp_path / "nowhere", project)
    assert r.code == 2 and "cannot build" in r.err and "LRB" in r.err


# ==== REAL: The owner's LRB folder ======================================================================

IN = (REAL_LRB / "in") if REAL_LRB else None
OUTS = (REAL_LRB / "Claude outputs") if REAL_LRB else None
LOGOS = (REAL_LRB / "files" / "Logos") if REAL_LRB else None
DINGTALK = (ROOT / "outputs" / REAL_ACCOUNT / "2026-11 月度计划" / "assets" / "superseded"
            / "钉钉原稿底稿 planstyle失败 Nov 2026.xlsx")


def colour(c):
    """A colour as a comparable value: its kind and its rgb, theme or indexed number."""
    if c is None:
        return None
    return (c.type, c.rgb if c.type == "rgb" else c.theme if c.type == "theme" else c.indexed)


def effective_widths(ws) -> dict[int, float]:
    """Width per column number. Excel saves neighbours of equal width as one grouped entry (H:J),
    which openpyxl lists under its first letter only, so a group is expanded before comparing."""
    from openpyxl.utils import column_index_from_string

    out = {}
    for letter, d in ws.column_dimensions.items():
        if d.width is None:
            continue
        first = d.min or column_index_from_string(letter)
        for idx in range(first, (d.max or first) + 1):
            out[idx] = d.width
    return out


def alignment(cell):
    a = cell.alignment
    return (a.horizontal or "general", a.vertical or "bottom", bool(a.wrap_text), a.indent or 0)


def edge(border, side):
    """One border side as a comparable value; a side that is not set at all compares as None."""
    s = getattr(border, side)
    return None if s is None or s.style is None else (s.style, colour(s.color))


def format_diffs(a: Path, b: Path, *, height_tol=0.01) -> list[str]:
    """Every format point that differs between two workbooks' first sheets. Values are not compared."""
    import openpyxl

    wa, wb_ = openpyxl.load_workbook(a).worksheets[0], openpyxl.load_workbook(b).worksheets[0]
    out: list[str] = []

    def cmp(label, x, y):
        if x != y:
            out.append(f"{label}: {x!r} != {y!r}")

    va, vb = wa.sheet_view, wb_.sheet_view
    # Excel writes some defaults out and leaves others out on a re-save (top-left cell A1, fit to 1x1 page,
    # bottom alignment, upper or lower case date format); a default compares equal to its absence.
    cmp("view.topLeftCell", va.topLeftCell or "A1", vb.topLeftCell or "A1")
    for attr in ("zoomScale", "showGridLines"):
        cmp(f"view.{attr}", getattr(va, attr), getattr(vb, attr))
    cmp("freeze", wa.freeze_panes, wb_.freeze_panes)
    cmp("selection", [(s.activeCell, s.sqref) for s in va.selection], [(s.activeCell, s.sqref) for s in vb.selection])
    for attr in ("orientation", "paperSize"):
        cmp(f"page.{attr}", getattr(wa.page_setup, attr), getattr(wb_.page_setup, attr))
    for attr in ("fitToWidth", "fitToHeight"):
        cmp(f"page.{attr}", getattr(wa.page_setup, attr) or 1, getattr(wb_.page_setup, attr) or 1)
    cmp("fitToPage", wa.sheet_properties.pageSetUpPr.fitToPage, wb_.sheet_properties.pageSetUpPr.fitToPage)
    cmp("hcentered", wa.print_options.horizontalCentered, wb_.print_options.horizontalCentered)
    cmp("margins", (wa.page_margins.left, wa.page_margins.right), (wb_.page_margins.left, wb_.page_margins.right))
    cmp("print titles", wa.print_title_rows, wb_.print_title_rows)
    cmp("merges", sorted(map(str, wa.merged_cells.ranges)), sorted(map(str, wb_.merged_cells.ranges)))
    cmp("sheet name", wa.title, wb_.title)
    widths_a, widths_b = effective_widths(wa), effective_widths(wb_)
    for col in sorted(set(widths_a) | set(widths_b)):
        cmp(f"col {col} width", widths_a.get(col), widths_b.get(col))
    rows = max(wa.max_row, wb_.max_row)
    for r in range(1, rows + 1):
        ha, hb = wa.row_dimensions[r].height, wb_.row_dimensions[r].height
        if (ha is None) != (hb is None) or (ha is not None and abs(ha - hb) > height_tol):
            out.append(f"row {r} height: {ha!r} != {hb!r}")
    for r in range(1, rows + 1):
        for c in range(1, max(wa.max_column, wb_.max_column) + 1):
            ca, cb = wa.cell(r, c), wb_.cell(r, c)
            if ca.value is None and cb.value is None:
                continue
            tag = ca.coordinate
            cmp(f"{tag} font", (ca.font.name, ca.font.sz, ca.font.b, ca.font.i, ca.font.u, colour(ca.font.color)),
                (cb.font.name, cb.font.sz, cb.font.b, cb.font.i, cb.font.u, colour(cb.font.color)))
            cmp(f"{tag} border", [edge(ca.border, s) for s in ("left", "right", "top", "bottom")],
                [edge(cb.border, s) for s in ("left", "right", "top", "bottom")])
            cmp(f"{tag} align", alignment(ca), alignment(cb))
            cmp(f"{tag} number format", ca.number_format.lower(), cb.number_format.lower())
            cmp(f"{tag} fill", (ca.fill.fill_type, colour(ca.fill.fgColor)), (cb.fill.fill_type, colour(cb.fill.fgColor)))
    cmp("image count", len(wa._images), len(wb_._images))
    for ia, ib in zip(wa._images, wb_._images):
        ea, eb = ia.anchor, ib.anchor
        cmp("logo anchor", (type(ea).__name__, ea._from.col, ea._from.row), (type(eb).__name__, eb._from.col, eb._from.row))
        if getattr(ea, "ext", None) is not None and getattr(eb, "ext", None) is not None:
            cmp("logo size", (round(ea.ext.width / 9525), round(ea.ext.height / 9525)),
                (round(eb.ext.width / 9525), round(eb.ext.height / 9525)))
    return out


@pytest.fixture(scope="module")
def real_project(tmp_path_factory):
    root = tmp_path_factory.mktemp("real-project")
    (root / "outputs" / REAL_ACCOUNT / "2026-11 月度计划").mkdir(parents=True)
    return root


@pytest.fixture(scope="module")
def real_plan(real_project):
    out = topic_dir(real_project) / f"{REAL_ACCOUNT} LRB Plan Nov 2026.xlsx"
    r = run("build", "--kind", "plan", "--source", str(IN / f"{REAL_ACCOUNT} LRB Monthly Plan.xlsx"),
            "--out", str(out), "--sheet", "Nov 2026", "--logo", str(LOGOS / f"{REAL_ACCOUNT} Logo.png"),
            "--lrb-dir", str(REAL_LRB), project=real_project)
    return r, out


@real
def test_real_plan_from_ros_in_copy_matches_his_reference_on_every_format_point_except_the_colours(real_plan):
    r, out = real_plan
    assert r.code == 0, r.out + r.err
    # The owner's reference carries the team draft's own colours; his ruling of 2026-10-17 ("No colours, per template")
    # takes them out of every new build, so fills are the one point where the build departs from that file
    diffs = format_diffs(out, OUTS / f"{REAL_ACCOUNT} LRB Plan Nov 2026.xlsx")
    assert not [d for d in diffs if " fill" not in d], "\n".join(diffs[:20])
    assert fill_cells(out) == {}


@real
def test_real_plan_lists_exactly_the_three_split_cells_and_the_fills_it_found(real_plan):
    r, out = real_plan
    assert [ln.split()[1].rstrip(":") for ln in r.tagged("KEPT-LINE")] == ["D3", "D8", "D10"], r.out
    assert {ln.split()[1].rstrip(":") for ln in r.tagged("OFFERED")} == {"D3", "D8", "D10"}
    assert not r.tagged("APPLIED") and not r.tagged("SHAPE")
    cleared = " ".join(r.tagged("CLEARED"))  # the owner's ruling 2026-10-17: the team draft's colours are removed
    assert "F3EEF8" in cleared and "B12" in cleared and "F9F6FC" in cleared, r.tagged("CLEARED")
    assert fill_cells(out) == {}
    done = cells(out)
    assert done["D3"] == TOPIC_PIPE_3 and done["D8"] == TOPIC_PIPE_8 and done["D10"] == TOPIC_PIPE_10


@real
def test_real_plan_passes_every_check(real_plan):
    r, _ = real_plan
    assert r.failed == set() and r.lines[-1] == "RESULT OK", r.out


@real
def test_real_dingtalk_copy_builds_without_error_and_has_the_same_layout(real_project, real_plan):
    if not DINGTALK.is_file():
        pytest.skip("the DingTalk copy is in the local outputs folder only")
    out = topic_dir(real_project) / "dingtalk.xlsx"
    r = run("build", "--kind", "plan", "--source", str(DINGTALK), "--out", str(out), "--sheet", "Nov 2026",
            "--logo", str(LOGOS / f"{REAL_ACCOUNT} Logo.png"), "--lrb-dir", str(REAL_LRB), project=real_project)
    assert r.code == 0, r.out + r.err
    shape = " ".join(r.tagged("SHAPE"))
    assert "column B to column A" in shape and "duplicate" in shape, r.out
    # same layout as the column-A copy's result; the draft's text differs, so the data rows may differ in height
    diffs = format_diffs(out, real_plan[1], height_tol=100)
    assert not diffs, "\n".join(diffs[:20])


@real
def test_real_check_flags_the_three_merged_cells_and_the_colours_in_ros_own_reference():
    r = check(OUTS / f"{REAL_ACCOUNT} LRB Plan Nov 2026.xlsx", REAL_LRB)
    assert r.failed == {"bilingual", "fills"}, r.out
    line = next(ln for ln in r.lines if ln.startswith("FAIL bilingual"))
    assert all(cell in line for cell in ("D3", "D8", "D10")), line
    fills = next(ln for ln in r.lines if ln.startswith("FAIL fills"))
    assert "F3EEF8" in fills and "B12" in fills and "F9F6FC" in fills, fills


@real
@pytest.mark.parametrize("name", [f"{a} LRB {kind}.xlsx" for a in REAL_ACCOUNTS
                                  for kind in ("Report Sep 2026", "Plan Nov 2026")])
def test_real_approved_files_carry_no_colours(name):
    path = OUTS / name
    if not path.is_file():
        pytest.skip(f"{name} is not in the owner's Claude outputs folder")
    assert fill_cells(path) == {}, "the standard says no fills; the approved files are what it was taken from"


@real
def test_real_report_from_ros_in_copy_matches_his_september_report(real_project):
    out = topic_dir(real_project) / f"{REAL_ACCOUNT} LRB Report Sep 2026.xlsx"
    r = run("build", "--kind", "report", "--source", str(IN / f"{REAL_ACCOUNT} LRB Monthly Report.xlsx"),
            "--out", str(out), "--sheet", "Sep 2026", "--logo", str(LOGOS / f"{REAL_ACCOUNT} Logo.png"),
            "--lrb-dir", str(REAL_LRB), project=real_project)
    assert r.code == 0, r.out + r.err
    assert "verify" in r.passed and "logo" in r.passed and "chars" in r.passed and "font" in r.passed, r.out
    # The only differences Excel's own re-save puts in the owner's file and the LibreOffice route does not, all
    # normalised by format_diffs: row heights a few hundredths of a point apart (59.05 against 59, allowed by
    # the 0.1 tolerance), neighbouring columns of equal width saved as one group, defaults written out or
    # left out, and the case of the date format. Nothing else may differ.
    diffs = format_diffs(out, OUTS / f"{REAL_ACCOUNT} LRB Report Sep 2026.xlsx", height_tol=0.1)
    assert not diffs, "\n".join(diffs[:20])


# ==== the skill's written text and its place in Easel ============================================


def _validator():
    spec = importlib.util.spec_from_file_location("validate_skills", ROOT / "scripts" / "validate_skills.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def test_skill_passes_the_contract_check_and_is_listed_once_under_plan():
    assert _validator().validate(SKILL_DIR / "SKILL.md") == []
    listed = [s for s in web.get_skills() if s["name"] == "lrb-monthly-excel"]
    assert len(listed) == 1 and listed[0]["layer"] == "plan" and listed[0]["description"].strip()


def test_skill_commands_match_the_script_options():
    spec = importlib.util.spec_from_file_location("validate_skill_commands", ROOT / "scripts" / "validate_skill_commands.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    mine = list(module.commands(SKILL_DIR / "SKILL.md"))
    assert mine, "SKILL.md documents no python command"
    supported, subcommands = module.cli_schema(SCRIPT)
    assert {"build", "check"} <= subcommands
    assert {line.split("lrb_excel.py", 1)[1].split()[0] for line, _ in mine} == {"build", "check"}
    for line, invocation in mine:
        assert invocation[0].endswith("skills/builtin/lrb-monthly-excel/scripts/lrb_excel.py"), line
        used = set(module.OPTION_RE.findall(line)) - {"--help"}
        assert used <= supported, (line, used - supported)


def test_skill_text_states_the_route_in_the_order_the_issue_gives():
    text = _text(SKILL_DIR / "SKILL.md")
    for needle in (
        "scripts/lrb_excel.py", "底稿", "--changes", "verify", "LRB Report and Plan Formatting Playbook",
        "planstyle.py", "reportstyle.py", "不要改", "底色",
    ):
        assert needle in text, f"SKILL.md does not say: {needle}"
    order = [text.index(k) for k in ("来源", "内容改动", "整理形状", "排版", "检查")]
    assert order == sorted(order), "the five steps are not in the order source, content, shape, formatting, checks"


def test_skill_text_says_colours_are_cleared_and_a_colour_left_in_the_file_fails_the_run():
    """The owner, 2026-10-17: "No colours, per template (Recommended)". The text must no longer say fills wait on him."""
    text = _text(SKILL_DIR / "SKILL.md")
    for needle in ("CLEARED", "`fills`", "底色"):
        assert needle in text, f"SKILL.md does not say: {needle}"
    for stale in ("只报告，不改", "要等负责人定", "不让运行失败", "FILLS"):
        assert stale not in text, f"SKILL.md still carries the old fills rule: {stale}"
    assert text.index("## 3.") < text.index("CLEARED") < text.index("## 4."), "the colour clearing belongs to the shape step"


def test_skill_text_says_how_a_change_names_its_cell_and_that_a_ruling_is_ros_quoted_words():
    """Issue 57 round 3: an Easel run could not apply a change whose text sits in several cells, and its list held summaries."""
    text = _text(SKILL_DIR / "SKILL.md")
    step2 = text[text.index("## 2."):text.index("## 3.")]
    for needle in ('"cell"', '"date"', '"column"', "原话", "引号"):
        assert needle in step2, f"step 2 does not say: {needle}"
    assert "C10" in step2 and "2026-11-27" in step2, "step 2 gives an example of each way to name the cell"


def test_content_calendar_hands_the_client_excel_to_the_route():
    text = _text(ROOT / "skills" / "builtin" / "skill-content-calendar" / "SKILL.md")
    assert "lrb-monthly-excel" in text
    assert "不要自己写" in text




# openpyxl is installed by Easel's own setup (ops/wsl-setup), not by `pip install -e .`, so CI may not have
# it. Every test that builds or reads a workbook says so and is skipped there; the written-text tests above
# do not need it and keep running everywhere.
_needs_openpyxl = pytest.mark.skipif(
    importlib.util.find_spec("openpyxl") is None,
    reason="openpyxl is not installed here (Easel's setup installs it into its own venv)",
)
for _name in [n for n in list(globals()) if n.startswith(
        ("test_check", "test_build", "test_spacer", "test_a_check", "test_a_lrb", "test_real"))]:
    globals()[_name] = _needs_openpyxl(globals()[_name])
